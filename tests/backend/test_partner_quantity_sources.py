"""Independent synthetic workbooks for quantity values and their source links."""
from datetime import datetime
import io
import zipfile

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from ekt.partner import read_partner


SKU = "000007_"


def _workbook(headers, rows, sheet_name):
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def _moq(rows, supplier="Systeme", sheet_name="MOQ rules"):
    if supplier == "IEK":
        headers = ["№", "Код 1С", "Артикул", "Наименование", "Мин. разр. к отгр."]
        records = [[n, sku, "SYNTHETIC", "Synthetic product", quantity]
                   for n, (sku, quantity) in enumerate(rows, 1)]
    else:
        headers = ["Наименование", "Другое", "Код", "Артикул", "Кратность"]
        records = [["Synthetic product", None, sku, "SYNTHETIC", quantity]
                   for sku, quantity in rows]
    return _workbook(headers, records, sheet_name)


def _monthly(rows, sheet_name="Monthly rules"):
    return _workbook(
        ["Наименование", "Код", "Артикул", "Кратность", datetime(2026, 8, 1)],
        [["Synthetic product", sku, "SYNTHETIC", quantity, 17] for sku, quantity in rows],
        sheet_name,
    )


def _import(files, supplier="Systeme", **options):
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as handle:
        for filename, payload in files:
            handle.writestr(filename, payload)
    return read_partner(archive.getvalue(), "synthetic.zip", supplier, "2026-09-22", **options)


def _product(bundle, sku=SKU):
    products = bundle["products"]
    selected = products.loc[products.sku_1c.eq(sku)]
    assert len(selected) == 1
    return selected.iloc[0]


def _assert_source(row, files, filename, source_row, field="order_multiple"):
    """Read the referenced cell independently instead of trusting the importer."""
    assert row.quantity_rule_source == f"{filename}:{source_row}"
    book = load_workbook(io.BytesIO(dict(files)[filename]), read_only=True, data_only=True)
    try:
        column = 5 if "moq" in filename.lower() else 4
        value = book.worksheets[0].cell(source_row, column).value
    finally:
        book.close()
    assert pd.notna(value)
    assert row[field] == value


@pytest.mark.parametrize("reverse_zip", [False, True])
@pytest.mark.parametrize("moq_first", [False, True])
def test_conflicting_moq_and_monthly_values_reject_any_file_order(reverse_zip, moq_first):
    moq_name = f"{'a' if moq_first else 'z'} MOQ.xlsx"
    monthly_name = f"{'z' if moq_first else 'a'} Ежемесячные продажи.xlsx"
    files = [(moq_name, _moq([(SKU, 10)])), (monthly_name, _monthly([(SKU, 5)]))]
    with pytest.raises(ValueError, match="Конфликт") as error:
        _import(files[::-1] if reverse_zip else files)
    message = str(error.value)
    for expected in ("Systeme", SKU, "order_multiple", "10", "5",
                     moq_name, monthly_name, "MOQ rules", "Monthly rules", ":2"):
        assert expected in message


@pytest.mark.parametrize("moq_first", [False, True])
def test_equal_values_keep_moq_source_regardless_of_file_order(moq_first):
    moq_name = f"{'a' if moq_first else 'z'} MOQ.xlsx"
    monthly_name = f"{'z' if moq_first else 'a'} Ежемесячные продажи.xlsx"
    files = [(moq_name, _moq([(SKU, 2.5)])), (monthly_name, _monthly([(SKU, 2.5)]))]
    row = _product(_import(files[::-1] if moq_first else files))
    assert row.order_multiple == 2.5
    _assert_source(row, files, moq_name, 2)


@pytest.mark.parametrize("moq_quantity,monthly_quantity,expected_source", [
    (None, 5, "monthly"),
    (10, None, "moq"),
    (None, None, None),
])
def test_blank_cells_do_not_create_or_replace_quantity_source(
    moq_quantity, monthly_quantity, expected_source,
):
    # Put the blank source last to catch a detached source overwrite.
    moq_first = moq_quantity is not None
    moq_name = f"{'a' if moq_first else 'z'} MOQ.xlsx"
    monthly_name = f"{'z' if moq_first else 'a'} Ежемесячные продажи.xlsx"
    files = [(moq_name, _moq([(SKU, moq_quantity)])),
             (monthly_name, _monthly([(SKU, monthly_quantity)]))]
    row = _product(_import(files))
    if expected_source is None:
        assert pd.isna(row.get("order_multiple"))
        assert pd.isna(row.get("quantity_rule_source"))
    else:
        name = moq_name if expected_source == "moq" else monthly_name
        _assert_source(row, files, name, 2)


def test_monthly_only_fractional_quantity_has_matching_source():
    quantity = 0.125
    files = [("Ежемесячные продажи.xlsx", _monthly([(SKU, quantity)]))]
    row = _product(_import(files))
    assert row.order_multiple == quantity
    _assert_source(row, files, files[0][0], 2)


def test_duplicate_moq_files_cannot_silently_replace_conflicting_values():
    files = [("a MOQ.xlsx", _moq([(SKU, 10)], sheet_name="First source")),
             ("z MOQ.xlsx", _moq([(SKU, 5)], sheet_name="Second source"))]
    with pytest.raises(ValueError, match="Конфликт") as error:
        _import(files)
    for expected in ("a MOQ.xlsx", "z MOQ.xlsx", "First source", "Second source",
                     "Systeme", SKU, "order_multiple", "10", "5"):
        assert expected in str(error.value)


def test_duplicate_moq_rows_cannot_silently_replace_conflicting_values():
    quantities = (10, 5)
    files = [("MOQ.xlsx", _moq([(SKU, quantity) for quantity in quantities]))]
    with pytest.raises(ValueError, match="Конфликт") as error:
        _import(files)
    for expected in ("MOQ.xlsx", "MOQ rules", ":2", ":3", "order_multiple", SKU):
        assert expected in str(error.value)


def test_equal_zero_sources_select_stable_file_and_row():
    quantity = 0
    files = [("z MOQ.xlsx", _moq([(SKU, quantity)])),
             ("a MOQ.xlsx", _moq([(SKU, quantity), (SKU, quantity)]))]
    row = _product(_import(files))
    _assert_source(row, files, "a MOQ.xlsx", 2)


@pytest.mark.parametrize("moq_quantity,monthly_quantity", [(0, 5), (5, 0), (1, 1.0000000001)])
def test_zero_and_close_but_distinct_values_are_real_conflicts(moq_quantity, monthly_quantity):
    files = [("MOQ.xlsx", _moq([(SKU, moq_quantity)])),
             ("Ежемесячные продажи.xlsx", _monthly([(SKU, monthly_quantity)]))]
    with pytest.raises(ValueError, match="Конфликт") as error:
        _import(files)
    for quantity in (moq_quantity, monthly_quantity):
        assert str(quantity) in str(error.value)


def test_rule_values_and_sources_are_independent_for_each_sku():
    files = [("MOQ.xlsx", _moq([(SKU, 10), ("000008_", None)])),
             ("Ежемесячные продажи.xlsx", _monthly([(SKU, None), ("000008_", 5)]))]
    bundle = _import(files)
    first, second = _product(bundle), _product(bundle, "000008_")
    assert (first.order_multiple, second.order_multiple) == (10, 5)
    _assert_source(first, files, "MOQ.xlsx", 2)
    _assert_source(second, files, "Ежемесячные продажи.xlsx", 3)


@pytest.mark.parametrize("meaning,field", [
    ("unknown", "unconfirmed_moq"),
    ("minimum", "min_order_qty"),
    ("multiple", "order_multiple"),
])
def test_iek_quantity_meaning_and_source_remain_explicit(meaning, field):
    files = [("MOQ IEK.xlsx", _moq([(SKU, 7)], supplier="IEK"))]
    row = _product(_import(files, supplier="IEK", iek_moq_meaning=meaning))
    assert row[field] == 7
    _assert_source(row, files, "MOQ IEK.xlsx", 2, field=field)
    for other in {"unconfirmed_moq", "min_order_qty", "order_multiple"} - {field}:
        assert pd.isna(row.get(other))
