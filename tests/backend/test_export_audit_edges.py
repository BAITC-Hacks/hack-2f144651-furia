import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from ekt.export import xlsx_bytes


@pytest.mark.parametrize("location", ["cell", "header", "metadata"])
def test_xlsx_rejects_text_that_exceeds_limit_after_control_escaping(location):
    value = "x" * 32764 + "\x0b"
    frame = pd.DataFrame({"name": [value if location == "cell" else "ok"]})
    metadata = {"note": value} if location == "metadata" else None
    if location == "header":
        frame.columns = [value]
    original = frame.copy(deep=True)
    sheet = "Calculation" if location == "metadata" else "Orders"

    with pytest.raises(ValueError, match=rf"XLSX: {sheet}.*32767"):
        xlsx_bytes(frame, metadata)

    pd.testing.assert_frame_equal(frame, original)


@pytest.mark.parametrize("index", [[7, 7], ["same", "same"], ["first", "second"]])
def test_xlsx_preserves_distinct_rows_with_nonstandard_indices(index):
    frame = pd.DataFrame({"sku_1c": ["00001\x0b", "00002\x0c"]}, index=index)
    original = frame.copy(deep=True)

    workbook = load_workbook(io.BytesIO(xlsx_bytes(frame)))

    assert workbook["Orders"]["A2"].value == r"00001\x0b"
    assert workbook["Orders"]["A3"].value == r"00002\x0c"
    pd.testing.assert_frame_equal(frame, original)


@pytest.mark.parametrize("character", ["\ud800", "\ufffe", "\uffff"])
@pytest.mark.parametrize("location", ["cell", "header", "metadata"])
def test_xlsx_rejects_unsupported_xml_unicode_with_cell_location(character, location):
    frame = pd.DataFrame({"name": [character if location == "cell" else "ok"]}, dtype=object)
    metadata = {"note": character} if location == "metadata" else None
    if location == "header":
        frame.columns = pd.Index([character], dtype=object)
    sheet = "Calculation" if location == "metadata" else "Orders"

    with pytest.raises(ValueError, match=rf"XLSX: {sheet}.*XML"):
        xlsx_bytes(frame, metadata)


def test_xlsx_preserves_duplicate_columns_when_escaped_headers_collide():
    frame = pd.DataFrame([["one", "two"]], columns=["name\x0b", r"name\x0b"])

    workbook = load_workbook(io.BytesIO(xlsx_bytes(frame)))

    assert list(workbook["Orders"].values) == [(r"name\x0b", r"name\x0b"), ("one", "two")]


def test_xlsx_round_trips_limit_text_unicode_and_formula_safe_sku_metadata():
    value = "x" * 32763 + "\x0b"
    frame = pd.DataFrame({"sku_1c": ["000001", "=1+1"], "name": [value, "Товар 🧰"]})
    metadata = {"=note": "=1+1", "details": {"text": "Товар 🧰"}}
    original = frame.copy(deep=True)

    workbook = load_workbook(io.BytesIO(xlsx_bytes(frame, metadata)))

    assert workbook["Orders"]["A2"].value == "000001"
    assert workbook["Orders"]["A3"].value == "'=1+1"
    assert workbook["Orders"]["A3"].data_type == "s"
    assert workbook["Orders"]["A3"].number_format == "@"
    assert workbook["Orders"]["B2"].value == "x" * 32763 + r"\x0b"
    assert workbook["Orders"]["B3"].value == "Товар 🧰"
    assert workbook["Calculation"]["A2"].value == "'=note"
    assert workbook["Calculation"]["B2"].value == "'=1+1"
    assert workbook["Calculation"]["B3"].value == '{"text": "Товар 🧰"}'
    pd.testing.assert_frame_equal(frame, original)
    assert metadata["=note"] == "=1+1"
