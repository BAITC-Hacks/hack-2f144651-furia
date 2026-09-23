"""Opt-in regression checks for private partner originals, never bundled in Git.

Run with EKT_PARTNER_ARCHIVE_DIR pointing to a directory containing IEK.zip and
Systeme.zip. The report date below is the explicit date of these test reports,
not an assertion that their balances are current today. Failures deliberately
omit source values, filenames, row numbers and exception details.
"""
from __future__ import annotations

from collections import defaultdict
from contextlib import contextmanager, redirect_stderr, redirect_stdout
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
import io
import math
import os
from pathlib import Path, PurePosixPath
import re
import zipfile

import pandas as pd
import pytest
from openpyxl import load_workbook

from ekt.partner import read_partner
from ekt.schema import validate


REPORT_DATE = pd.Timestamp("2026-09-22")
MONTH_NAMES = (
    "янв", "фев", "мар", "апр", "май", "июн",
    "июл", "авг", "сен", "окт", "ноя", "дек",
)


def _check(condition, supplier, table, check):
    # No rewritten assert: pytest must never display a private row/value diff.
    if not bool(condition):
        pytest.fail(f"{supplier}/{table}/{check}", pytrace=False)


@contextmanager
def _private_failure(supplier, table):
    try:
        yield
    except Exception:
        raise pytest.fail.Exception(f"{supplier}/{table}/unexpected_error", pytrace=False) from None


def _number(value):
    if value is None or str(value).strip() == "":
        return None
    return Decimal(str(value).replace("\xa0", "").replace(" ", "").replace(",", "."))


def _same_number(actual, source):
    expected = _number(source)
    if expected is None:
        return bool(pd.isna(actual))
    return pd.notna(actual) and math.isclose(float(actual), float(expected), rel_tol=1e-12, abs_tol=1e-9)


def _code(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip() or None


def _source_date(value):
    if isinstance(value, (datetime, date)):
        return pd.Timestamp(value).normalize()
    # Partner text dates use DD.MM.YYYY; do not let pandas guess month-first.
    for pattern in ("%d.%m.%Y", "%d.%m.%Y %H:%M:%S"):
        try:
            return pd.Timestamp(datetime.strptime(str(value).strip(), pattern)).normalize()
        except ValueError:
            continue
    raise ValueError("Unsupported source date format")


def _month(value):
    if isinstance(value, (datetime, date)):
        return pd.Timestamp(value.year, value.month, 1)
    label = str(value or "").lower()
    year = re.search(r"\b20\d{2}\b", label)
    for number, stem in enumerate(MONTH_NAMES, 1):
        if year and stem in label:
            return pd.Timestamp(int(year[0]), number, 1)
    return None


def _kind(filename):
    name = filename.lower()
    if "динамик" in name:
        return "sales"
    if "ежемесяч" in name:
        return "historical_stock" if "остат" in name else "monthly_sales"
    if "moq" in name:
        return "moq"
    if "сезон" in name:
        return "seasonal_prior"
    if "путь" in name or "пути" in name:
        return "inbound"
    return "unknown"


def _sku_column(kind, supplier):
    return {
        "sales": 3,
        "monthly_sales": 1,
        "historical_stock": 2,
        "moq": 1 if supplier == "IEK" else 2,
        "inbound": 0 if supplier == "IEK" else 2,
    }[kind]


def _sample(frame):
    """At most three positional samples per distinct source and snapshot kind."""
    if frame.empty:
        return frame
    columns = ["source_file", "source_sheet"]
    if "snapshot_kind" in frame:
        columns.append("snapshot_kind")
    pieces = []
    for _, group in frame.groupby(columns, dropna=False, sort=False):
        pieces.append(group.iloc[sorted({0, len(group) // 2, len(group) - 1})])
    return pd.concat(pieces)


@dataclass(repr=False)
class OriginalCase:
    supplier: str
    bundle: object
    members: dict = field(default_factory=dict)
    samples: dict = field(default_factory=dict)
    cells: dict = field(default_factory=dict)
    sheets: dict = field(default_factory=dict)
    signed_sales: dict = field(default_factory=dict)
    moq_samples: object = None
    stock_payload: bytes = b""
    moq_payload: bytes = b""

    def __repr__(self):
        return f"OriginalCase({self.supplier})"

    def raw(self, row):
        return self.cells[(row.source_file, row.source_sheet, int(row.source_row))]


@pytest.fixture(scope="module", params=("IEK", "Systeme"))
def original(request):
    directory = os.environ.get("EKT_PARTNER_ARCHIVE_DIR")
    if not directory:
        pytest.skip("Private originals omitted; set EKT_PARTNER_ARCHIVE_DIR to opt in")
    supplier = request.param
    with _private_failure(supplier, "fixture"):
        archive_path = Path(directory) / f"{supplier}.zip"
        _check(archive_path.is_file(), supplier, "archive", "missing_private_archive")
        data = archive_path.read_bytes()
        # Import exactly once per supplier/module. Suppress any accidental library
        # output, including in a failing run; private data stays in memory.
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            bundle = read_partner(
                data, archive_path.name, supplier, REPORT_DATE,
                confirmed_scope=None, iek_moq_meaning="unknown", inbound_base_units=False,
            )
        case = OriginalCase(supplier, bundle)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for member in archive.infolist():
                name = PurePosixPath(member.filename.replace("\\", "/")).name
                if not member.is_dir() and name.lower().endswith(".xlsx") and not name.startswith("~$"):
                    _check(name not in case.members, supplier, "archive", "duplicate_filename")
                    case.members[name] = archive.read(member)
        wanted = defaultdict(lambda: defaultdict(set))
        for table in ("products", "sales", "monthly_sales", "stock_snapshots", "inbound", "seasonal_prior"):
            frame = bundle[table]
            case.samples[table] = frame if table == "seasonal_prior" else _sample(frame)
            for row in case.samples[table].itertuples():
                wanted[row.source_file][row.source_sheet].update((1, 2, int(row.source_row)))
        products = bundle["products"]
        _check("quantity_rule_source" in products, supplier, "products", "moq_provenance")
        case.moq_samples = _sample(products.loc[products.quantity_rule_source.notna()])
        moq_rows = defaultdict(set)
        for row in case.moq_samples.itertuples():
            filename, source_row = row.quantity_rule_source.rsplit(":", 1)
            moq_rows[filename].add(int(source_row))
        # Each workbook is opened once for independent verification, read-only.
        # Only requested rows plus a pair of signed sales samples are cached.
        for filename, payload in case.members.items():
            if _kind(filename) == "historical_stock":
                case.stock_payload = payload
            elif _kind(filename) == "moq":
                case.moq_payload = payload
            with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
                book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            try:
                case.sheets[filename] = tuple(book.sheetnames)
                if filename in moq_rows:
                    wanted[filename][book.worksheets[0].title].update({1, *moq_rows[filename]})
                for sheet_name, row_numbers in wanted[filename].items():
                    sheet = book[sheet_name]
                    for number, values in enumerate(sheet.iter_rows(max_row=max(row_numbers), values_only=True), 1):
                        if number in row_numbers:
                            case.cells[(filename, sheet_name, number)] = values
                        if _kind(filename) == "sales" and number > 1 and values[0] is not None and values[3] is not None:
                            quantity = _number(values[7])
                            if quantity is None or quantity != 0:
                                sign = "missing" if quantity is None else "negative" if quantity < 0 else "positive"
                                key = (filename, sheet_name, sign)
                                if key not in case.signed_sales:
                                    case.signed_sales[key] = number
                                    case.cells[(filename, sheet_name, number)] = values
            finally:
                book.close()
        # Drop duplicate XLSX bytes after caching the small verification sample.
        case.members = {name: _kind(name) for name in case.members}
    return case


def test_original_inventory_and_conservative_contract(original):
    case, supplier = original, original.supplier
    with _private_failure(supplier, "contract"):
        bundle = case.bundle
        _check(len(case.members) == 6, supplier, "archive", "six_xlsx")
        _check(set(case.members.values()) == {"sales", "monthly_sales", "historical_stock", "moq", "inbound", "seasonal_prior"}, supplier, "archive", "report_types")
        _check(any("XLSX: 6" in note for note in bundle.notes), supplier, "archive", "adapter_inventory")
        _check(bundle.mode == "partner", supplier, "bundle", "data_mode")
        for table, frame in bundle.tables.items():
            if frame.empty:
                continue
            _check(frame.supplier_id.eq(supplier).all(), supplier, table, "supplier")
            _check(frame.data_mode.eq("partner").all(), supplier, table, "data_mode")
            _check(frame.source_file.isin(case.members).all(), supplier, table, "source_file")
            _check(frame.source_sheet.notna().all() and frame.source_row.str.fullmatch(r"[1-9]\d*").fillna(False).all(), supplier, table, "provenance")
        for table in ("products", "sales", "monthly_sales", "stock_snapshots", "inbound", "seasonal_prior"):
            _check(not bundle[table].empty, supplier, table, "nonempty")
        _check(bundle["sales"].customer_id.isna().all(), supplier, "sales", "no_invented_customers")
        for table in ("stockouts", "policies", "growth_plan"):
            _check(bundle[table].empty, supplier, table, "no_invented_inputs")
        for table in ("monthly_sales", "stock_snapshots", "inbound"):
            _check(bundle[table].warehouse_scope.eq("UNSPECIFIED").all(), supplier, table, "unconfirmed_scope")
        inbound = bundle["inbound"]
        _check(inbound.status.eq("pending").all(), supplier, "inbound", "unconfirmed_status")
        _check(inbound.qty_base_unit.isna().all(), supplier, "inbound", "unconfirmed_base_units")
        _check(inbound.qty_source_unit.notna().all(), supplier, "inbound", "source_quantity_preserved")
        if supplier == "IEK":
            _check(bundle["stock_snapshots"].snapshot_kind.eq("month_start").all(), supplier, "stock_snapshots", "no_current_balance")
            for column in ("min_order_qty", "order_multiple"):
                _check(column not in bundle["products"] or bundle["products"][column].isna().all(), supplier, "products", "moq_meaning_unknown")


def test_original_product_and_moq_samples(original):
    case, supplier = original, original.supplier
    with _private_failure(supplier, "products"):
        for row in case.samples["products"].itertuples():
            source = case.raw(row)
            column = _sku_column(case.members[row.source_file], supplier)
            _check(isinstance(row.sku_1c, str) and row.sku_1c == _code(source[column]), supplier, "products", "source_sku")
        _check(not case.moq_samples.empty, supplier, "products", "moq_samples")
        filename = next(name for name, kind in case.members.items() if kind == "moq")
        isolated = read_partner(case.moq_payload, filename, supplier, REPORT_DATE)
        products = isolated["products"].set_index("sku_1c")
        field = "unconfirmed_moq" if supplier == "IEK" else "order_multiple"
        for row in case.moq_samples.itertuples():
            filename, number = row.quantity_rule_source.rsplit(":", 1)
            source = case.cells[(filename, case.sheets[filename][0], int(number))]
            _check(row.sku_1c == _code(source[_sku_column("moq", supplier)]), supplier, "products", "moq_source_sku")
            _check(_same_number(products.loc[row.sku_1c, field], source[4]), supplier, "products", "standalone_moq_source_quantity")


def test_original_merged_moq_provenance(original, request):
    case, supplier = original, original.supplier
    if supplier == "Systeme":
        # DATA/D3: monthly order_multiple overwrites MOQ while retaining the
        # MOQ source reference. Keep this defect visible until its owner fixes it.
        request.applymarker(pytest.mark.xfail(
            strict=True, reason="DATA/D3: merged Systeme MOQ provenance is stale",
        ))
    with _private_failure(supplier, "products"):
        field = "unconfirmed_moq" if supplier == "IEK" else "order_multiple"
        for row in case.moq_samples.itertuples():
            filename, number = row.quantity_rule_source.rsplit(":", 1)
            source = case.cells[(filename, case.sheets[filename][0], int(number))]
            _check(_same_number(getattr(row, field), source[4]), supplier, "products", "merged_moq_source_quantity")


def test_original_sales_source_samples_and_signs(original):
    case, supplier = original, original.supplier
    with _private_failure(supplier, "sales"):
        frame = case.bundle["sales"]
        rows = list(case.samples["sales"].itertuples())
        _check({"positive", "negative"}.issubset({key[2] for key in case.signed_sales}), supplier, "sales", "source_sign_coverage")
        for (filename, sheet, _), number in case.signed_sales.items():
            selected = frame.loc[frame.source_file.eq(filename) & frame.source_sheet.eq(sheet) & frame.source_row.eq(str(number))]
            _check(len(selected) == 1, supplier, "sales", "signed_source_row_preserved")
            rows.extend(selected.itertuples())
        for row in rows:
            source = case.raw(row)
            _check(row.sku_1c == _code(source[3]), supplier, "sales", "source_sku")
            _check(_same_number(row.quantity_signed, source[7]), supplier, "sales", "signed_quantity")
            _check(row.date == _source_date(source[0]), supplier, "sales", "source_date")
            _check(row.unit == _code(source[5]), supplier, "sales", "source_unit")
            _check(row.warehouse_scope == (_code(source[6]) or "UNSPECIFIED"), supplier, "sales", "source_scope")
            document = str(source[2] or "").lower()
            kind = next((kind for prefix, kind in (("расходная накладная", "sale"), ("приходная накладная", "receipt"), ("заказ покупателя", "customer_order"), ("возврат", "return")) if document.startswith(prefix)), "unknown")
            _check(row.document_type == kind, supplier, "sales", "document_semantics")
        # An unknown source quantity remains unknown and is a validation issue;
        # it must not be silently filled with zero or dropped to make data pass.
        errors = validate(case.bundle)
        _check(frame.quantity_signed.isna().any() == any(error.startswith("sales.quantity_signed:") for error in errors), supplier, "sales", "missing_quantity_validation")


def test_original_monthly_source_samples(original):
    case, supplier = original, original.supplier
    with _private_failure(supplier, "monthly_sales"):
        for row in case.samples["monthly_sales"].itertuples():
            source = case.raw(row)
            kind = case.members[row.source_file]
            header_row = 2 if kind == "inbound" else 1
            headers = case.cells[(row.source_file, row.source_sheet, header_row)]
            columns = [i for i, value in enumerate(headers) if _month(value) == row.month]
            _check(len(columns) == 1, supplier, "monthly_sales", "source_month")
            _check(row.sku_1c == _code(source[_sku_column(kind, supplier)]), supplier, "monthly_sales", "source_sku")
            _check(_same_number(row.qty_net, source[columns[0]]), supplier, "monthly_sales", "source_quantity")
            end = row.month + pd.offsets.MonthEnd(0)
            _check(row.coverage_start == row.month and row.coverage_end == min(end, REPORT_DATE), supplier, "monthly_sales", "coverage")
            _check(bool(row.is_complete) == (end <= REPORT_DATE), supplier, "monthly_sales", "partial_report_month")


def test_original_stock_source_samples(original):
    case, supplier = original, original.supplier
    with _private_failure(supplier, "stock_snapshots"):
        for row in case.samples["stock_snapshots"].itertuples():
            source = case.raw(row)
            _check(row.sku_1c == _code(source[2]), supplier, "stock_snapshots", "source_sku")
            if row.snapshot_kind == "current":
                _check(supplier == "Systeme" and row.as_of == REPORT_DATE, supplier, "stock_snapshots", "explicit_report_date")
                for field, column in (("on_hand", 49), ("reserved", 50), ("available", 51)):
                    _check(_same_number(getattr(row, field), source[column]), supplier, "stock_snapshots", f"source_{field}")
            else:
                headers = case.cells[(row.source_file, row.source_sheet, 1)]
                columns = [i for i, value in enumerate(headers) if _month(value) == row.as_of]
                _check(len(columns) == 1, supplier, "stock_snapshots", "source_month")
                _check(_same_number(row.on_hand, source[columns[0]]), supplier, "stock_snapshots", "source_quantity")
                _check(pd.isna(row.available) and pd.isna(row.reserved), supplier, "stock_snapshots", "no_invented_available_or_reserve")
        if supplier == "Systeme":
            _check(case.bundle["stock_snapshots"].snapshot_kind.eq("current").any(), supplier, "stock_snapshots", "current_report_present")


def test_original_inbound_source_samples(original):
    case, supplier = original, original.supplier
    with _private_failure(supplier, "inbound"):
        for row in case.samples["inbound"].itertuples():
            source = case.raw(row)
            _check(row.sku_1c == _code(source[_sku_column("inbound", supplier)]), supplier, "inbound", "source_sku")
            column = int(row.order_id.rsplit(":", 1)[1]) - 1 if supplier == "IEK" else 54
            _check(_same_number(row.qty_source_unit, source[column]), supplier, "inbound", "source_quantity")
            if supplier == "IEK":
                headers = case.cells[(row.source_file, row.source_sheet, 1)]
                _check(row.eta_header == _code(headers[column]), supplier, "inbound", "source_eta_header")


def test_original_stock_product_names_are_not_ordinals(original):
    case, supplier = original, original.supplier
    with _private_failure(supplier, "products"):
        filename = next(name for name, kind in case.members.items() if kind == "historical_stock")
        # Isolate the report so later imports cannot conceal an incorrect name.
        isolated = read_partner(case.stock_payload, filename, supplier, REPORT_DATE)
        products = isolated["products"].set_index("sku_1c")
        samples = case.samples["stock_snapshots"]
        samples = samples.loc[samples.source_file.eq(filename)]
        _check(not samples.empty, supplier, "products", "stock_name_samples")
        for row in samples.itertuples():
            headers = case.cells[(row.source_file, row.source_sheet, 1)]
            name_column = 0 if supplier == "IEK" else 1
            _check(headers[name_column] == "Номенклатура", supplier, "products", "stock_name_header")
            if supplier == "Systeme":
                _check(headers[0] == "№", supplier, "products", "numbered_stock_layout")
            expected = _code(case.raw(row)[name_column])
            _check(products.loc[row.sku_1c, "name"] == expected, supplier, "products", "stock_name_from_labelled_cell")


def test_original_seasonal_source_samples(original):
    case, supplier = original, original.supplier
    with _private_failure(supplier, "seasonal_prior"):
        frame = case.bundle["seasonal_prior"]
        _check(len(frame) == 12 and set(frame.month_of_year) == set(range(1, 13)), supplier, "seasonal_prior", "twelve_months")
        column = 5 if supplier == "IEK" else 11
        raw_factors = [float(_number(case.raw(row)[column])) for row in frame.itertuples()]
        mean = sum(raw_factors) / 12
        for row, factor in zip(frame.itertuples(), raw_factors):
            _check(_same_number(row.raw_factor, factor), supplier, "seasonal_prior", "raw_factor")
            _check(_same_number(row.factor, factor / mean), supplier, "seasonal_prior", "normalized_factor")
            _check(row.known_as_of == REPORT_DATE, supplier, "seasonal_prior", "explicit_report_date")
