import io
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from ekt.application import import_canonical_files
from ekt.demo import demo_bundle
from ekt.export import xlsx_bytes
from ekt.ingest import ImportBudget, read_canonical
from ekt.partner import read_partner


HEADERS = ["supplier_id", "sku_1c", "name", "category_id", "unit"]
TEXT_ROWS = [
    ["NA", "000001", "NULL", "A", "pcs"],
    ["NULL", "NA", "N/A", "B", "pcs"],
    ["N/A", "NULL", "NA", "C", "pcs"],
    ["SUP", "N/A", "", "D", "pcs"],
]


def _xlsx(rows, sheet="products"):
    output = io.BytesIO()
    workbook = Workbook()
    page = workbook.active
    page.title = sheet
    page.append(HEADERS)
    for row in rows:
        page.append(row)
    workbook.save(output)
    return output.getvalue()


def _csv(rows):
    lines = [",".join(HEADERS)]
    lines.extend(",".join(row) for row in rows)
    return ("\n".join(lines) + "\n").encode()


@pytest.mark.parametrize("filename,payload", [
    ("products.csv", _csv(TEXT_ROWS)),
    ("products.xlsx", _xlsx(TEXT_ROWS)),
])
def test_text_codes_and_real_blanks_are_equal_in_csv_and_xlsx(filename, payload):
    frame = read_canonical(payload, filename)["products"]
    assert frame["supplier_id"].tolist() == ["NA", "NULL", "N/A", "SUP"]
    assert frame["sku_1c"].tolist() == ["000001", "NA", "NULL", "N/A"]
    assert frame["name"].iloc[:3].tolist() == ["NULL", "N/A", "NA"]
    assert pd.isna(frame["name"].iloc[3])


def test_xlsx_export_escapes_xml_controls_only_in_export_copy_and_keeps_text_sku():
    source = pd.DataFrame({"sku_1c": ["Synthetic\x0bitem"], "name": ["=1+1"]})
    original = source.copy(deep=True)
    exported = load_workbook(io.BytesIO(xlsx_bytes(source, {"note": "meta\x0bdata"})))
    assert exported["Orders"]["A2"].value == r"Synthetic\x0bitem"
    assert exported["Orders"]["A2"].data_type == "s"
    assert exported["Orders"]["B2"].value == "'=1+1"
    assert exported["Calculation"]["B2"].value == r"meta\x0bdata"
    pd.testing.assert_frame_equal(source, original)


def test_xlsx_export_sanitizes_untrusted_header_text_in_export_copy():
    source = pd.DataFrame({"=name\x0b": [1]})
    workbook = load_workbook(io.BytesIO(xlsx_bytes(source)))
    header = workbook["Orders"]["A1"]
    assert header.value == "'=name\\x0b"
    assert header.data_type == "s"
    assert source.columns.tolist() == ["=name\x0b"]


@pytest.mark.parametrize("frame,metadata,sheet", [
    (pd.DataFrame({"name": ["x" * 32768]}), None, "Orders"),
    (pd.DataFrame({"x": [1]}).rename(columns={"x": "h" * 32768}), None, "Orders"),
    (pd.DataFrame({"name": ["ok"]}), {"note": "x" * 32768}, "Calculation"),
])
def test_xlsx_rejects_unrepresentable_long_text_with_location(frame, metadata, sheet):
    with pytest.raises(ValueError, match=rf"{sheet}.*32767"):
        xlsx_bytes(frame, metadata)


def test_cumulative_import_budget_rejects_batch_before_previous_changes():
    previous = demo_bundle()
    original = previous.copy()
    first = b"supplier_id,sku_1c,name\nS,000001,one\n"
    second = b"supplier_id,sku_1c,name\nS,000002,two\n"
    budget = ImportBudget(max_bytes=len(first) + len(second) - 1)
    with pytest.raises(ValueError, match="входных байтов"):
        import_canonical_files(
            [("products.csv", first), ("products.csv", second)],
            previous=previous,
            budget=budget,
        )
    for name in previous.tables:
        pd.testing.assert_frame_equal(previous[name], original[name])


def test_cumulative_row_and_cell_budgets_are_enforced_on_small_fixture():
    payload = _csv(TEXT_ROWS)
    with pytest.raises(ValueError, match="строк"):
        read_canonical(payload, "products.csv", budget=ImportBudget(max_rows=2))
    with pytest.raises(ValueError, match="ячеек"):
        read_canonical(payload, "products.csv", budget=ImportBudget(max_cells=5))


def test_nested_xlsx_expanded_size_is_charged_before_member_read():
    workbook = _xlsx(TEXT_ROWS)
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("products.xlsx", workbook)
    budget = ImportBudget(max_expanded_bytes=100)
    with pytest.raises(ValueError, match="распакованных байтов"):
        read_canonical(archive.getvalue(), "bundle.zip", budget=budget)


def test_partner_import_enforces_workbook_row_budget_before_report_parsing():
    with pytest.raises(ValueError, match=r"unknown\.xlsx.*строк"):
        read_partner(
            _xlsx(TEXT_ROWS), "unknown.xlsx", "IEK", "2026-09-22",
            budget=ImportBudget(max_rows=2),
        )


def test_partner_zip_file_count_is_checked_before_member_payloads_are_read():
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("a.xlsx", _xlsx([]))
        package.writestr("b.xlsx", _xlsx([]))
    with pytest.raises(ValueError, match="файлов"):
        read_partner(
            archive.getvalue(), "reports.zip", "IEK", "2026-09-22",
            budget=ImportBudget(max_files=1),
        )


def test_nested_partner_xlsx_xml_size_is_charged_before_workbook_parse():
    workbook = _xlsx([])
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("unknown.xlsx", workbook)
    budget = ImportBudget(max_expanded_bytes=len(workbook) + 10)
    with pytest.raises(ValueError, match="распакованных байтов"):
        read_partner(
            archive.getvalue(), "reports.zip", "IEK", "2026-09-22", budget=budget,
        )
