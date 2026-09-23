"""Small resource limits prove rejection before workbook/dataframe allocation."""
import csv
import io
import re
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook

from ekt import ingest, partner
from ekt.application import import_canonical_files
from ekt.demo import demo_bundle
from ekt.ingest import ImportBudget, read_canonical


def workbook_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = "products"
    sheet.append(["supplier_id", "sku_1c", "name"])
    sheet.append(["NA", "000001", "NULL"])
    result = io.BytesIO()
    book.save(result)
    book.close()
    return result.getvalue()


def rewrite_sheet(payload, transform):
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(payload)) as source, zipfile.ZipFile(result, "w") as target:
        for entry in source.infolist():
            content = source.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = transform(content.decode()).encode()
            target.writestr(entry, content)
    return result.getvalue()


def pack(entries):
    result = io.BytesIO()
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
        for filename, data in entries:
            archive.writestr(filename, data)
    return result.getvalue()


def unexpected_parser(*args, **kwargs):
    pytest.fail("budget must reject before the expensive parser")


def test_csv_long_text_does_not_inherit_or_modify_global_field_limit():
    original_limit = csv.field_size_limit()
    long_text = "x" * (original_limit + 1)
    frame = read_canonical(
        f"supplier_id,sku_1c,name\nNA,000001,{long_text}\n".encode(), "products.csv",
    )["products"]
    assert frame.iloc[0]["name"] == long_text
    assert frame.iloc[0]["supplier_id"] == "NA"
    assert frame.iloc[0]["sku_1c"] == "000001"
    assert csv.field_size_limit() == original_limit


@pytest.mark.parametrize("text", [
    "", "\n", "\r\n", "a,b\r\n\r\nc,d", '"a\nb","c"\n',
    '"a""b",c\n', '"a"b"c,d\n', 'a"b,c\n', '"",\n',
    '"a\r\nb",c\r\nd,e\r\n', "a,b\rc,d\r", '"a,b,c',
])
def test_csv_budget_counts_records_with_quotes_and_newlines(text):
    expected = list(csv.reader(io.StringIO(text, newline="")))
    budget = ImportBudget()
    ingest._count_csv(text, ",", budget)
    assert budget.rows == len(expected)
    assert budget.cells == len(expected) * max(map(len, expected), default=0)


def test_csv_cell_limit_rejects_before_pandas(monkeypatch):
    monkeypatch.setattr(ingest.pd, "read_csv", unexpected_parser)
    with pytest.raises(ValueError, match="ячеек"):
        read_canonical(b'supplier_id,sku_1c,name\nNA,000001,"line1\nline2"\n',
                       "products.csv", budget=ImportBudget(max_cells=5))


def test_csv_short_records_count_padding_before_pandas(monkeypatch):
    monkeypatch.setattr(ingest.pd, "read_csv", unexpected_parser)
    with pytest.raises(ValueError, match="ячеек"):
        read_canonical(b"supplier_id,sku_1c,name,unit,category_id\nNA\nNULL\nN/A\n",
                       "products.csv", budget=ImportBudget(max_cells=12))


@pytest.mark.parametrize("transform,error,budget", [
    (lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:ZZZ2"', xml),
     "ячеек", {"max_cells": 20}),
    (lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:C100"', xml),
     "строк", {"max_rows": 5}),
    (lambda xml: re.sub(r'<dimension[^>]*/>', "", xml).replace('r="B2"', 'r="ZZZ2"'),
     "ячеек", {"max_cells": 20}),
    (lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:A1"', xml)
     .replace('r="B2"', 'r="ZZZ2"'), "ячеек", {"max_cells": 20}),
])
@pytest.mark.parametrize("kind", ["canonical", "partner"])
def test_worksheet_bounds_reject_before_any_excel_parser(monkeypatch, transform, error, budget, kind):
    payload = rewrite_sheet(workbook_bytes(), transform)
    monkeypatch.setattr(ingest.pd, "ExcelFile", unexpected_parser)
    monkeypatch.setattr(partner, "load_workbook", unexpected_parser)
    with pytest.raises(ValueError, match=error):
        if kind == "canonical":
            read_canonical(payload, "products.xlsx", budget=ImportBudget(**budget))
        else:
            partner.read_partner(payload, "unknown.xlsx", "IEK", "2026-09-22", budget=ImportBudget(**budget))


def test_workbook_row_budget_is_cumulative_across_files_and_preserves_previous():
    previous = demo_bundle()
    original = previous.copy()
    files = [("products.xlsx", workbook_bytes())] * 2
    with pytest.raises(ValueError, match="строк"):
        import_canonical_files(files, previous=previous, budget=ImportBudget(max_rows=3))
    for name in previous.tables:
        pd.testing.assert_frame_equal(previous[name], original[name])


@pytest.mark.parametrize("target,filename", [
    ("/xl/custom/sheet.bin", "xl/custom/sheet.bin"),
    ("../custom/sheet.xml", "custom/sheet.xml"),
])
def test_worksheet_relationship_targets_are_counted_even_outside_default_folder(monkeypatch, target, filename):
    result = io.BytesIO()
    original_path = "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(io.BytesIO(workbook_bytes())) as source, zipfile.ZipFile(result, "w") as output:
        for entry in source.infolist():
            content = source.read(entry)
            if entry.filename == "xl/_rels/workbook.xml.rels":
                content = content.replace(("/" + original_path).encode(), target.encode())
            output.writestr(filename if entry.filename == original_path else entry.filename, content)
    payload = result.getvalue()
    # The relocated file is a real, readable workbook, not a malformed fixture.
    assert read_canonical(payload, "products.xlsx")["products"].sku_1c.tolist() == ["000001"]
    monkeypatch.setattr(ingest.pd, "ExcelFile", unexpected_parser)
    with pytest.raises(ValueError, match="ячеек"):
        read_canonical(payload, "products.xlsx", budget=ImportBudget(max_cells=5))


def test_nested_workbooks_charge_each_xml_package_before_loading(monkeypatch):
    payload = workbook_bytes()
    with zipfile.ZipFile(io.BytesIO(payload)) as workbook:
        expanded = sum(entry.file_size for entry in workbook.infolist())
    calls = []
    real_load = partner.load_workbook

    def record_load(*args, **kwargs):
        calls.append(1)
        return real_load(*args, **kwargs)

    monkeypatch.setattr(partner, "load_workbook", record_load)
    archive = pack([("one.xlsx", payload), ("two.xlsx", payload)])
    with pytest.raises(ValueError, match="распакованных байтов"):
        partner.read_partner(archive, "reports.zip", "IEK", "2026-09-22",
                             budget=ImportBudget(max_expanded_bytes=2 * len(payload) + expanded))
    assert len(calls) == 1


def test_partner_zip_order_uses_basenames_and_reads_one_payload_at_a_time(monkeypatch):
    archive = pack([("a/z.xlsx", b"z"), ("z/a.xlsx", b"a")])
    real_read = zipfile.ZipFile.read
    reads = []

    def record_read(self, member, *args, **kwargs):
        reads.append(member.filename)
        return real_read(self, member, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "read", record_read)
    members = partner.xlsx_members(archive, "reports.zip")
    try:
        assert next(members) == ("a.xlsx", b"a")
        assert reads == ["z/a.xlsx"]
        assert next(members) == ("z.xlsx", b"z")
        assert reads == ["z/a.xlsx", "a/z.xlsx"]
    finally:
        members.close()


def test_partner_closes_workbook_after_adapter_error(monkeypatch):
    real_load = partner.load_workbook
    closed = []

    def record_load(*args, **kwargs):
        book = real_load(*args, **kwargs)
        real_close = book.close

        def close():
            closed.append(True)
            real_close()

        book.close = close
        return book

    monkeypatch.setattr(partner, "load_workbook", record_load)
    with pytest.raises(ValueError, match="заголовки динамики"):
        partner.read_partner(workbook_bytes(), "Динамика продаж.xlsx", "IEK", "2026-09-22")
    assert closed == [True]
