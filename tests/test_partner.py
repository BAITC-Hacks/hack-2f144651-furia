"""Synthetic workbook fixtures follow the published audit; not real partner data."""
from datetime import datetime
import io
import zipfile
import pandas as pd
import pytest
from openpyxl import Workbook
from ekt.partner import read_partner
from ekt.engine import calculate
from ekt.schema import SCHEMAS, normalize, validate


def workbook_bytes(book):
    target = io.BytesIO()
    book.save(target)
    return target.getvalue()


def bundle_archive(supplier):
    systeme = supplier == "Systeme"
    files = {}
    months = pd.date_range("2024-01-01", "2026-09-01", freq="MS")
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист_1"
    sheet.append(["Дата", "Номер", "Документ", "Код", "Наименование", "Ед.", "Склад", "Количество"])
    sheet.append([datetime(2024, 1, 1), "DOC-1", "Расходная накладная 1", "000001_", "Synthetic only", "шт", "Склад A", 105])
    sheet.append([datetime(2024, 1, 2), "DOC-2", "Расходная накладная 2", "000001_", "Synthetic only", "шт", "Склад A", -5])
    sheet.append([datetime(2024, 1, 3), "DOC-3", "Приходная накладная", "000001_", "Synthetic only", "шт", "Склад A", 200])
    sheet.append([datetime(2024, 1, 4), "DOC-4", "Заказ покупателя", "000001_", "Synthetic only", "шт", "Склад A", 1000])
    sheet.append([None, "Итого", None, None, None, None, None, 1300])
    files["Динамика продаж.xlsx"] = workbook_bytes(book)
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист_1"
    prefix = ["Наименование", "Код", "Артикул", "Кратность"] if systeme else ["Наименование", "Код"]
    data_prefix = ["Synthetic only", "000001_", "A000", 6] if systeme else ["Synthetic only", "000001_"]
    sheet.append(prefix + list(months.to_pydatetime()) + ["Итого"])
    sheet.append([None])
    sheet.append(data_prefix + [100] * len(months) + [3300])
    files["Ежемесячные продажи.xlsx"] = workbook_bytes(book)
    book = Workbook()
    sheet = book.active
    sheet.title = "Лист_1"
    prefix = ["Наименование", "Артикул", "Код", "Ед."] if systeme else ["Наименование", "Артикул", "Код"]
    data_prefix = ["Synthetic only", "A000", "000001_", "шт"] if systeme else ["Synthetic only", "A000", "000001_"]
    sheet.append(prefix + list(months.to_pydatetime()) + ["Итого"])
    sheet.append([None])
    sheet.append([None, None, None, "нач. остаток"] if not systeme else [None])
    sheet.append(data_prefix + [25] * len(months) + [825])
    files["Ежемесячные остатки.xlsx"] = workbook_bytes(book)
    book = Workbook()
    sheet = book.active
    if systeme:
        sheet.append(["Имя", "Другое", "Код", "Артикул", "Кратность"])
        sheet.append([None])
        sheet.append(["Synthetic only", None, "000001_", "A000", 6])
    else:
        sheet.title = "Лист7"
        sheet.append(["N", "Код 1С", "Артикул", "Наименование", "Мин. разр. к отгр."])
        sheet.append([1, "000001_", "A000", "Synthetic only", 10])
    files["MOQ.xlsx"] = workbook_bytes(book)
    book = Workbook()
    sheet = book.active
    russian_months = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]
    start, col = (11, 12) if systeme else (28, 6)
    sheet.title = "Лист1" if systeme else "Сезонность"
    for n, name in enumerate(russian_months, start):
        sheet.cell(n, 2, name)
        sheet.cell(n, col, 2 if name == "июнь" else 1)
    files["Сезонность.xlsx"] = workbook_bytes(book)
    book = Workbook()
    sheet = book.active
    if systeme:
        sheet.title = "TDSheet"
        for col, label in [(2, "Артикул"), (3, "Код"), (4, "Наименование"), (5, "Категория 2026"), (50, "Остаток"), (51, "Резерв"), (52, "Свободный остаток"), (55, "СЭ в пути 24.09")]:
            sheet.cell(2, col, label)
        for col, value in [(2, "A000"), (3, "000001_"), (4, "Synthetic only"), (5, "1"), (50, 100), (51, 20), (52, 80), (55, 30)]:
            sheet.cell(3, col, value)
        for col, month in enumerate(months, 7):
            sheet.cell(2, col, month.to_pydatetime())
            sheet.cell(3, col, 999)
        sheet.cell(2, 42, "Последние 12 мес")
        sheet.cell(3, 42, "=SUM(AC3:AO3)")
        sheet.cell(3, 43, "=AP3/12")
        sheet.cell(3, 44, -.5)
        sheet.cell(3, 45, -.8)
        files["Товар в пути_SystemElectric.xlsx"] = workbook_bytes(book)
    else:
        sheet.title = "Лист4"
        sheet.append(["Код 1С", "Артикул", "Наименование", "до 24.09.2026", "25.09.2026", "01.10.2026", "01.11.2026", "неизвестно", "02.12.2026"])
        sheet.append(["000001_", "A000", "Synthetic only", 30, 0, 0, 0, 0, 0])
        files["Путь ИЭК 22.09.2026.xlsx"] = workbook_bytes(book)
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as handle:
        for name, data in files.items():
            handle.writestr(name, data)
    return archive.getvalue()


@pytest.mark.parametrize("supplier", ["IEK", "Systeme"])
def test_all_six_audited_reports_import_without_inventing_fields(supplier):
    bundle = read_partner(bundle_archive(supplier), "partner.zip", supplier, "2026-09-22")
    assert len(bundle["products"]) == 1
    assert bundle["products"].sku_1c.iloc[0] == "000001_"
    assert bundle["sales"].customer_id.isna().all()
    assert bundle["sales"].quantity_signed.tolist() == [105, -5, 200, 1000]
    assert len(bundle["monthly_sales"]) == 33
    assert bundle["monthly_sales"].qty_net.eq(100).all()
    assert not bundle["monthly_sales"].iloc[-1].is_complete
    assert bundle["stockouts"].empty
    assert bundle["policies"].empty
    assert bundle["growth_plan"].empty
    assert bundle["inbound"].status.eq("pending").all()
    assert bundle["inbound"].qty_base_unit.isna().all()
    assert bundle["seasonal_prior"].factor.mean() == pytest.approx(1)
    assert len(bundle["seasonal_prior"]) == 12
    assert any("XLSX: 6" in note for note in bundle.notes)
    assert not validate(bundle)
    if supplier == "IEK":
        assert bundle["stock_snapshots"].snapshot_kind.eq("month_start").all()
        assert bundle["products"].unconfirmed_moq.iloc[0] == 10
    else:
        current = bundle["stock_snapshots"].query("snapshot_kind == 'current'").iloc[0]
        assert current.available == 80
        assert current.reserved == 20
        assert current.as_of == pd.Timestamp("2026-09-22")
        assert bundle["products"].order_multiple.iloc[0] == 6


def test_systeme_with_explicit_scope_units_and_policy_calculates():
    bundle = read_partner(bundle_archive("Systeme"), "partner.zip", "Systeme", "2026-09-22", "POOL_CONFIRMED", inbound_base_units=True)
    bundle.tables["policies"] = pd.DataFrame([["Systeme", "1", 14, 7, 3, 0, 1, "manual"]], columns=SCHEMAS["policies"])
    result = calculate(bundle, "2026-09-22").rows.iloc[0]
    assert pd.notna(result.recommended_qty)
    assert result.available_stock == 80
    assert result.inbound_within_horizon == 30
    assert result.excluded_oneoff_qty == 0
    assert result.recommended_qty % 6 == 0


def test_iek_historical_stock_never_becomes_current():
    bundle = read_partner(bundle_archive("IEK"), "partner.zip", "IEK", "2026-09-22", "POOL", "minimum", True)
    bundle.tables["policies"] = pd.DataFrame([["IEK", None, 14, 7, 3, 0, 1, "manual"]], columns=SCHEMAS["policies"])
    row = calculate(bundle, "2026-09-22").rows.iloc[0]
    assert pd.isna(row.recommended_qty)
    assert "текущего остатка" in row.explanation


def test_invalid_headers_are_rejected():
    book = Workbook()
    book.active.append(["Неожиданный заголовок"])
    with pytest.raises(ValueError, match="заголовки"):
        read_partner(workbook_bytes(book), "Динамика продаж.xlsx", "IEK", "2026-09-22")
