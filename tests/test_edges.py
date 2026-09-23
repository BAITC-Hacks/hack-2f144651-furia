import io
import zipfile
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from ekt.demo import demo_bundle, canonical_zip
from ekt.schema import SCHEMAS, normalize, validate
from ekt.engine import calculate, order_quantity
from ekt.ingest import read_canonical, parse_csv
from ekt.export import csv_bytes, xlsx_bytes
from ekt.review import initial_edits, approve


def test_no_double_count_of_monthly_and_invoices():
    bundle = demo_bundle()
    first = calculate(bundle, "2026-09-22")
    bundle.tables["monthly_sales"] = bundle["monthly_sales"].iloc[0:0]
    second = calculate(bundle, "2026-09-22")
    np.testing.assert_allclose(first.rows.expected_demand_horizon, second.rows.expected_demand_horizon)


def test_free_stock_is_not_reduced_by_reserve_twice(make_bundle):
    bundle = make_bundle()
    bundle["stock_snapshots"].loc[0, ["on_hand", "reserved", "available"]] = [100, 20, 80]
    assert calculate(bundle, "2026-08-31").rows.iloc[0].available_stock == 80


@pytest.mark.parametrize("mutation, message", [("missing_stock", "остат"), ("historical_stock", "остат"), ("missing_lead", "полит"), ("unspecified", "склад"), ("units", "Единиц")])
def test_missing_inputs_block_precise_order(make_bundle, mutation, message):
    bundle = make_bundle()
    if mutation == "missing_stock":
        bundle["stock_snapshots"].loc[0, ["available", "on_hand", "reserved"]] = np.nan
    elif mutation == "historical_stock":
        bundle["stock_snapshots"]["snapshot_kind"] = "month_start"
    elif mutation == "missing_lead":
        bundle.tables["policies"] = bundle["policies"].iloc[0:0]
    elif mutation == "unspecified":
        for table in bundle.tables.values():
            if "warehouse_scope" in table:
                table["warehouse_scope"] = "UNSPECIFIED"
    else:
        bundle["sales"]["unit"] = "упак"
    row = calculate(bundle, "2026-08-31").rows.iloc[0]
    assert pd.isna(row.recommended_qty)
    assert message.lower() in row.explanation.lower()


def test_late_unknown_overdue_and_cancelled_inbound_do_not_reduce_need(make_bundle):
    bundle = make_bundle()
    baseline = calculate(bundle, "2026-08-31").rows.iloc[0].recommended_qty
    records = [["S", "000001_", "W", "late", 1000, "2026-12-01", "expected", "confirmed"],
               ["S", "000001_", "W", "unknown", 1000, None, "unknown", "confirmed"],
               ["S", "000001_", "W", "old", 1000, "2026-08-01", "expected", "confirmed"],
               ["S", "000001_", "W", "cancel", 1000, "2026-09-01", "expected", "cancelled"]]
    bundle.tables["inbound"] = pd.DataFrame(records, columns=SCHEMAS["inbound"])
    row = calculate(bundle, "2026-08-31").rows.iloc[0]
    assert row.recommended_qty == baseline
    assert row.inbound_late == row.inbound_unknown_eta == row.inbound_overdue == 1000
    assert row.inbound_within_horizon == 0


def test_duplicate_party_detected_but_same_order_other_sku_allowed(make_bundle):
    bundle = make_bundle()
    item = ["S", "000001_", "W", "in-1", 50, "2026-09-03", "expected", "confirmed"]
    bundle.tables["inbound"] = pd.DataFrame([item, item], columns=SCHEMAS["inbound"])
    with pytest.raises(ValueError, match="дубли"):
        calculate(bundle, "2026-08-31")
    product = bundle["products"].copy()
    product["sku_1c"] = "000002"
    bundle.tables["products"] = pd.concat([bundle["products"], product], ignore_index=True)
    bundle["inbound"].loc[1, "sku_1c"] = "000002"
    assert not validate(normalize(bundle))


def test_future_sales_do_not_leak_into_prediction(make_bundle):
    bundle = make_bundle()
    before = calculate(bundle, "2026-08-31").rows.iloc[0]
    future = bundle["sales"].iloc[[0]].copy()
    future["date"] = pd.Timestamp("2026-09-15")
    future["quantity_signed"] = 1e9
    bundle.tables["sales"] = pd.concat([bundle["sales"], future], ignore_index=True)
    after = calculate(bundle, "2026-08-31").rows.iloc[0]
    assert before.expected_demand_horizon == after.expected_demand_horizon


def test_identifiers_and_missing_values_survive_import():
    frame = parse_csv(b"supplier_id,sku_1c,available\nS,000001_,\nS,000002,0\n")
    assert frame.sku_1c.tolist() == ["000001_", "000002"]
    assert pd.isna(frame.available.iloc[0])
    assert frame.available.iloc[1] == "0"


def test_formula_injection_text_escaped_numbers_kept_numeric():
    frame = pd.DataFrame({"sku_1c": ["000001_"], "name": [" =HYPERLINK(\"bad\")"], "final_qty": [5.5], "balance": [-3.]})
    book = load_workbook(io.BytesIO(xlsx_bytes(frame)))
    assert book["Orders"]["B2"].data_type == "s"
    assert book["Orders"]["B2"].value.startswith("'")
    assert book["Orders"]["C2"].value == 5.5
    assert book["Orders"]["D2"].value == -3
    assert "' =HYPERLINK" in csv_bytes(frame).decode("utf-8-sig")


def test_invalid_multiple_finite_values_and_manual_reason(make_bundle):
    with pytest.raises(ValueError):
        order_quantity(10, 0, 0, 0, 0)
    bundle = make_bundle()
    bundle["sales"].loc[0, "quantity_signed"] = np.inf
    with pytest.raises(ValueError, match="Infinity"):
        calculate(bundle, "2026-08-31")
    result = calculate(make_bundle(), "2026-08-31")
    edits = initial_edits(result.rows)
    edits.loc[0, "adjusted_qty"] = 0
    with pytest.raises(ValueError, match="причину"):
        approve(result, edits)


def test_archive_paths_and_mode_are_checked():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("../products.csv", "supplier_id,sku_1c\nS,1\n")
    with pytest.raises(ValueError, match="путь"):
        read_canonical(data.getvalue(), "input.zip")
    with pytest.raises(ValueError, match="Синтетический"):
        read_canonical(canonical_zip(demo_bundle()), "demo.zip", "partner")
