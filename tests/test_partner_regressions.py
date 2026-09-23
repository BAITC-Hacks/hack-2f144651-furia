from datetime import date, datetime
import io
import zipfile

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from ekt.partner import eta_value, numeric, read_partner
from test_partner import bundle_archive, workbook_bytes


def members(supplier="Systeme"):
    with zipfile.ZipFile(io.BytesIO(bundle_archive(supplier))) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def pack(files):
    result = io.BytesIO()
    with zipfile.ZipFile(result, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return result.getvalue()


def change_cell(payload, row, column, value):
    book = load_workbook(io.BytesIO(payload))
    book.active.cell(row, column).value = value
    result = workbook_bytes(book)
    book.close()
    return result


def test_d1_primary_report_wins_with_explicit_conflict_in_any_zip_order():
    files = members()
    a = read_partner(pack(files), "synthetic.zip", "Systeme", "2026-09-22")
    b = read_partner(pack(dict(reversed(list(files.items())))), "synthetic.zip", "Systeme", "2026-09-22")
    pd.testing.assert_frame_equal(a["monthly_sales"], b["monthly_sales"])
    assert a.notes == b.notes
    assert a["monthly_sales"].qty_net.eq(100).all()
    conflicts = [n for n in a.notes if "Расхождение" in n]
    assert len(conflicts) == 33
    assert all("100" in n and "999" in n and "000001_" in n for n in conflicts)
    assert "2024-01" in conflicts[0]
    assert "Ежемесячные продажи.xlsx" in conflicts[0]
    assert "Товар в пути_SystemElectric.xlsx" in conflicts[0]


def test_d1_two_conflicting_primary_reports_are_not_overwritten():
    files = members()
    files["Ежемесячные продажи копия.xlsx"] = change_cell(files["Ежемесячные продажи.xlsx"], 3, 5, 123)
    for order in (files, dict(reversed(list(files.items())))):
        with pytest.raises(ValueError, match="Конфликт основных.*000001_.*2024-01"):
            read_partner(pack(order), "synthetic.zip", "Systeme", "2026-09-22")


def test_d1_equal_primary_reports_do_not_duplicate_sales():
    files = members()
    files["Ежемесячные продажи копия.xlsx"] = files["Ежемесячные продажи.xlsx"]
    result = read_partner(pack(files), "synthetic.zip", "Systeme", "2026-09-22")
    assert len(result["monthly_sales"]) == 33
    assert result["monthly_sales"].qty_net.sum() == 3300


@pytest.mark.parametrize("value", [datetime(2026, 9, 24, 15), date(2026, 9, 24), pd.Timestamp("2026-09-24"), "2026-09-24", "до 24.09.2026", "24/09/2026", "24.09"])
def test_d2_supported_eta_formats(value):
    assert eta_value(value, pd.Timestamp("2026-09-22")) == pd.Timestamp("2026-09-24")


@pytest.mark.parametrize("value", [None, pd.NaT, "", "неизвестно", "31.02.2026", "2026-02-31", "24.09 - 01.10", "24.09.26", 46289])
def test_d2_unknown_or_invalid_eta_remains_unknown(value):
    assert eta_value(value, pd.Timestamp("2026-09-22")) is None


def test_d2_year_rollover_is_not_guessed():
    assert eta_value("05.01", pd.Timestamp("2026-12-22")) is None
    assert eta_value("05.01.2026", pd.Timestamp("2026-12-22")) == pd.Timestamp("2026-01-05")


def test_d2_unresolved_inbound_header_is_visible():
    files = members()
    name = "Товар в пути_SystemElectric.xlsx"
    files[name] = change_cell(files[name], 2, 55, "СЭ в пути 05.01")
    result = read_partner(pack(files), "synthetic.zip", "Systeme", "2026-12-22", "W", inbound_base_units=True)
    assert result["inbound"].eta.isna().all()
    assert result["inbound"].qty_base_unit.iloc[0] == 30
    assert any("05.01" in n and "ETA" in n and "год" in n for n in result.notes)


@pytest.mark.parametrize("value", ["not a number", "NaN", np.nan, np.inf, "1e999", True])
def test_d3_bad_numeric_cells_do_not_become_missing(value):
    with pytest.raises(ValueError, match="числ"):
        numeric(value)


def test_d3_empty_and_localized_numbers_are_preserved():
    assert numeric(None) is None
    assert numeric("  ") is None
    assert numeric("1 234,5") == 1234.5
    assert numeric(-5) == -5


def test_d3_bad_sales_date_does_not_silently_drop_sku():
    files = members("IEK")
    files["Динамика продаж.xlsx"] = change_cell(files["Динамика продаж.xlsx"], 2, 1, "31.02.2026")
    with pytest.raises(ValueError, match="Динамика продаж.*2.*дат"):
        read_partner(pack(files), "synthetic.zip", "IEK", "2026-09-22")


def test_d3_bad_monthly_quantity_has_file_context():
    files = members()
    files["Ежемесячные продажи.xlsx"] = change_cell(files["Ежемесячные продажи.xlsx"], 3, 5, "bad")
    with pytest.raises(ValueError, match="Ежемесячные продажи.*числ"):
        read_partner(pack(files), "synthetic.zip", "Systeme", "2026-09-22")
