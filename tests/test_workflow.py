import io
from pathlib import Path
from openpyxl import load_workbook
import pandas as pd
from streamlit.testing.v1 import AppTest
from ekt.demo import demo_bundle, canonical_zip
from ekt.ingest import read_canonical
from ekt.engine import calculate, order_quantity
from ekt.review import initial_edits, approve, export_frame
from ekt.export import csv_bytes, xlsx_bytes


def test_t1_t2_t3_order_arithmetic():
    assert order_quantity(100, 20, 30, 25, 10) == (65, 70)
    assert order_quantity(100, 20, 30, 45, 10) == (45, 50)
    assert order_quantity(100, 20, 200, 25, 10, 50) == (0, 0)
    assert order_quantity(11, 0, 0, 0, 6, 10) == (11, 12)


def test_load_calculate_explain_adjust_approve_export():
    source = demo_bundle()
    bundle = read_canonical(canonical_zip(source), "demo.zip", "synthetic")
    result = calculate(bundle, "2026-09-22")
    assert result.rows.recommended_qty.notna().all()
    assert result.rows.explanation.str.contains("прогноз").all()
    edits = initial_edits(result.rows)
    edits.loc[0, "adjusted_qty"] = 0.0
    edits.loc[0, "reason"] = "Проверка ручного нуля"
    edits.loc[0, "selected"] = True
    approval = approve(result, edits)
    frame = export_frame(result, edits, approval)
    assert frame.loc[0, "final_qty"] == 0
    assert frame.approval_status.eq("approved").all()
    csv = csv_bytes(frame)
    assert csv.startswith(b"\xef\xbb\xbf")
    parsed = pd.read_csv(io.BytesIO(csv), sep=";", dtype={"sku_1c": "string"})
    assert parsed.loc[0, "sku_1c"] == "000001_"
    assert parsed.loc[0, "final_qty"] == 0
    book = load_workbook(io.BytesIO(xlsx_bytes(frame)), data_only=False)
    headers = [c.value for c in book["Orders"][1]]
    sku_cell = book["Orders"].cell(2, headers.index("sku_1c") + 1)
    assert sku_cell.value == "000001_" and sku_cell.data_type == "s"
    edits.loc[0, "adjusted_qty"] = 1
    assert export_frame(result, edits, approval).approval_status.eq("draft").all()
    revised = calculate(bundle, "2026-09-22", remove_oneoffs=False)
    assert export_frame(revised, edits, approval).approval_status.eq("draft").all()


def test_streamlit_full_demo():
    app = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=30).run()
    assert not app.exception
    app.button[0].click().run()
    assert not app.exception
    next(b for b in app.button if b.label == "Рассчитать предложения").click().run()
    assert not app.exception
    assert app.metric[1].value == "4"
    next(b for b in app.button if b.label == "Утвердить выбранные позиции").click().run()
    assert not app.exception
    assert any("утверждены" in message.value for message in app.success)
    app.toggle[0].set_value(False).run()
    assert "calculation" not in app.session_state
    assert "approval" not in app.session_state
    assert not app.exception
